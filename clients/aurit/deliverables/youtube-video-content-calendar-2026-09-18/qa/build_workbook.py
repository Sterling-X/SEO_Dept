#!/usr/bin/env python3
"""Build the Aurit YouTube video content calendar workbook from the workflow's structured draft.
Usage: build_workbook.py draft.json legal.json out.xlsx
Format lineage: Kalish six-tab structure + Hicks Legend/Order/Internal Link Target conventions and header style."""
import json, sys, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

NAVY = "101F3C"; FONT = "Arial"
HDR_FONT = Font(name=FONT, size=10, bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor=NAVY)
BODY = Font(name=FONT, size=10)
TITLE = Font(name=FONT, size=14, bold=True, color=NAVY)
SUB = Font(name=FONT, size=10, italic=True, color="444444")
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Side(style="thin", color="D9D9D9"); BORDER = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)

def clean(v):
    if v is None: return ""
    if isinstance(v, list): v = "\n".join(f"• {clean(x)}" if not str(x).startswith("•") else clean(x) for x in v)
    s = str(v).replace("—", "-").replace("–", "-").replace(" ", " ")
    return re.sub(r"[ \t]+", " ", s).strip()

def sheet(wb, name, title, intro_lines, headers, rows, widths, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = name
    ws["A1"] = title; ws["A1"].font = TITLE
    r = 2
    for line in intro_lines:
        ws.cell(row=r, column=1, value=clean(line)).font = SUB; ws.cell(row=r, column=1).alignment = Alignment(wrap_text=False); r += 1
    hdr_row = r + 1
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=hdr_row, column=c, value=h); cell.font = HDR_FONT; cell.fill = HDR_FILL; cell.alignment = WRAP; cell.border = BORDER
    for i, row in enumerate(rows, hdr_row + 1):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=i, column=c, value=clean(v)); cell.font = BODY; cell.alignment = WRAP; cell.border = BORDER
    for c, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = ws.cell(row=hdr_row + 1, column=1)
    ws.row_dimensions[hdr_row].height = 30
    return ws

def main(draft_path, legal_path, out):
    d = json.load(open(draft_path)); legal = json.load(open(legal_path)) if legal_path != "-" else {"results": [], "could_not_check": []}
    wb = Workbook()
    L = d["legend"]
    ws = wb.active; ws.title = "Legend"
    ws["A1"] = "Aurit Mediation: YouTube Video Content Calendar"; ws["A1"].font = TITLE
    legend_rows = [("Client", "The Aurit Center for Divorce Mediation (auritmediation.com), Arizona divorce and family mediation"),
                   ("Prepared", "September 18, 2026 by Rocket Clicks SEO"),
                   ("Purpose", L["purpose"]), ("How to use this workbook", L["how_to_use"]), ("Column notes", L["column_notes"]),
                   ("Scope and evidence", L["scope_and_evidence"]), ("Verification summary", d.get("verification_summary", "")),
                   ("Open questions", L["open_questions"]), ("Could not check", d.get("could_not_check", [])),
                   ("Tabs", ["Calendar: ordered production sequence by track and batch", "Recording Briefs: one per recording, with legal claims to verify", "Shorts: clip plan per recording plus FAQ shorts", "Topic Research: every candidate that reached verification, with decision and evidence", "Publishing Plan: production, distribution, governance, measurement", "AI Search Briefs: answer packaging for AI search per recording", "Legal Verification: live check of each Arizona legal claim in the briefs", "Evidence Sources: files and sweeps behind the recommendations"])]
    for i, (k, v) in enumerate(legend_rows, 3):
        a = ws.cell(row=i, column=1, value=k); a.font = Font(name=FONT, size=10, bold=True); a.alignment = WRAP
        b = ws.cell(row=i, column=2, value=clean(v)); b.font = BODY; b.alignment = WRAP
    ws.column_dimensions["A"].width = 26; ws.column_dimensions["B"].width = 120

    cal = sorted(d["calendar"], key=lambda r: r["order"])
    sheet(wb, "Calendar", "Calendar: recording and publishing sequence",
          ["Ordered production sequence grouped by track and batch. Order is the suggested sequence, not a hard deadline; assign dates to suit mediator availability.",
           f"{sum(1 for r in cal if r['batch'].lower().startswith('track 0') or 'activate' in r['batch'].lower())} existing-asset activations and {sum(1 for r in cal if not (r['batch'].lower().startswith('track 0') or 'activate' in r['batch'].lower()))} new recordings. See Topic Research for evidence and decisions, Recording Briefs for scripts, Legal Verification before recording.",
           "Internal Link Target is the live auritmediation.com page the video supports and links to; 'ARCHITECTURE GATE' in Notes means a new page would need approval first."],
          ["Order", "Track / Batch", "Recording Title", "Format", "Minutes", "Priority", "Hub Section", "Target Query / Search Intent", "Topic Signal", "Why It Matters Commercially", "Internal Link Target", "Featured Speaker", "CTA / Next Step", "Existing Asset", "Status", "Notes"],
          [[r["order"], r["batch"], r["recording_title"], r["format"], r["minutes"], r["priority"], r["hub_section"], r["target_query"], r["topic_signal"], r["why_it_matters_commercially"], r["internal_link_target"], r["featured_speaker"], r["cta_next_step"], r["existing_asset"], r["status"], r["notes"]] for r in cal],
          [7, 20, 42, 16, 9, 11, 24, 32, 30, 44, 38, 22, 30, 28, 11, 40])

    sheet(wb, "Recording Briefs", "Recording briefs and website mapping",
          ["Use one clear answer, one fictional example (labelled as such) and one topic-matched next step in each recording.",
           "Existing pages are recommended starting points under the architecture gate; add a new page only for a distinct, approved intent.",
           "The presenting mediator must verify every item in 'Legal Claims to Verify' against the Legal Verification tab and current Arizona authority before recording."],
          ["Recording Title", "Target Search Question", "Intended Viewer", "Opening Hook", "Talking Points", "Example or Visual", "Existing Page to Support", "Service Destination", "Consultation CTA", "Script Review Focus", "Legal Claims to Verify"],
          [[b["recording_title"], b["target_search_question"], b["intended_viewer"], b["opening_hook"], b["talking_points"], b["example_or_visual"], b["existing_page_to_support"], b["service_destination"], b["consultation_cta"], b["script_review_focus"], b["legal_claims_to_verify"]] for b in d["recording_briefs"]],
          [40, 32, 30, 36, 60, 36, 38, 34, 32, 40, 44])

    sheet(wb, "Shorts", "Short-form release calendar",
          ["Clips are separate edits of approved footage, not additional recordings. Keep the full answer's qualifications in each cut.",
           "Distribution includes YouTube Shorts, Instagram Reels, Facebook and Google Business Profile video posts (Aurit already publishes weekly GBP posts)."],
          ["Clip Title / Opening Question", "Source Video", "Length", "Distribution", "Next Step", "Status", "Editor", "Published URL"],
          [[s["clip_title"], s["source_video"], s["length"], s["distribution"], s["next_step"], s["status"], "", ""] for s in d["shorts"]],
          [50, 44, 12, 34, 40, 11, 14, 24])

    sheet(wb, "Topic Research", "Topic research and priority decisions",
          ["Every candidate that reached adversarial verification, including those excluded, with the evidence type, period and decision.",
           "Figures appear only where the evidence contains them; autocomplete and forum evidence show recurrence, not volume."],
          ["Candidate Topic", "Priority", "Evidence Type", "Recording Decision", "Consultation / Business Rationale", "What the Evidence Supports", "Evidence Period", "Sources", "Production Implication"],
          [[t["candidate_topic"], t["priority"], t["evidence_type"], t["recording_decision"], t["business_rationale"], t["what_the_evidence_supports"], t["evidence_period"], t["sources"], t["production_implication"]] for t in d["topic_research"]],
          [36, 11, 26, 24, 44, 50, 20, 44, 44])

    sheet(wb, "Publishing Plan", "Production, distribution, governance and measurement",
          ["Proposed owners are roles to assign, not confirmed team members.",
           "Prioritisation is provisional until YouTube Studio, Search Console, GA4, call-tracking and intake baselines are captured."],
          ["Area", "Action", "Why / Definition", "Proposed Owner", "When"],
          [[p["area"], p["action"], p["why_or_definition"], p["proposed_owner"], p["when"]] for p in d["publishing_plan"]],
          [24, 60, 60, 24, 24])

    sheet(wb, "AI Search Briefs", "Answer briefs for AI search and helpful video discovery",
          ["Proposed prompts to test, not observed AI query volumes.",
           "For each asset: answer promptly, explain the factors, show one useful example, cite the primary source, and route to the topic-matched next step."],
          ["Recording Title", "Proposed Client Prompt", "Direct-Answer Focus", "Related Questions to Cover", "Original Supporting Asset", "Source to Review / Cite", "Existing Destination", "Topic-Matched Next Step"],
          [[a["recording_title"], a["proposed_client_prompt"], a["direct_answer_focus"], a["related_questions_to_cover"], a["original_supporting_asset"], a["source_to_review_or_cite"], a["existing_destination"], a["topic_matched_next_step"]] for a in d["ai_search_briefs"]],
          [40, 44, 40, 40, 32, 40, 38, 36])

    sheet(wb, "Legal Verification", "Live verification of Arizona legal claims used in the briefs (September 18, 2026)",
          ["Each claim was checked against current primary authority fetched live on the date shown. 'Not verified' means the authority could not be retrieved, not that the claim is wrong.",
           "This tab supports, and does not replace, review by the presenting attorney-mediator before recording."] + ([f"Could not check: {'; '.join(legal.get('could_not_check', []))}"] if legal.get("could_not_check") else []),
          ["Recording Title", "Claim", "Status", "Authority", "Authority URL", "Accessed", "Finding", "Correction for Brief"],
          [[r["recording_title"], r["claim"], r["status"], r["authority"], r["authority_url"], r["accessed"], r["finding"], r["correction_for_brief"]] for r in legal.get("results", [])],
          [36, 44, 14, 30, 40, 12, 50, 44])

    ev = [["aurit_brief.md", "Consolidated verified client facts, assets, architecture gate, format lineage", "Compiled 2026-09-18 from auritmediation.com fetches, ClickUp, Drive"],
          ["aurit_keyword_research.csv / aurit_kw_vol.csv / aurit_priorities.csv", "Client keyword research (103 terms, volumes, SERP features incl. Video and AI Overview) and content priorities", "Drive: 'Aurit | Comp. Gap, KW Research & Nav Layout', modified 2026-07-30"],
          ["aurit_ai_content_plan.csv", "Planned AI-first supporting articles with primary conversational queries", "Drive: 'Aurit_Detailed_AI_Supporting_Content_Plan_2026_27', modified 2026-09-08"],
          ["aurit_live_url_match*.csv", "Client hub architecture maps: H2 sections, intents, best live URL, coverage", "Drive: 'Aurit_Live_URL_Match_Production', modified 2026-09-10"],
          ["aurit_live_pages.csv", "192 live direct-200 pages with title, H1, word count, template", "Rocket Clicks parity audit crawl, 2026-09-17"],
          ["autocomplete_sweep.csv", "1,012 unique Google and YouTube autocomplete suggestions (40 seeds x 12 prefixes)", "Collected 2026-09-18"],
          ["youtube_search_sweep.csv", "585 YouTube results across 30 queries with views, age, length, channel", "Collected 2026-09-18"],
          ["Client 2026 Strategic Plan, Q3-4 tab rows 85-88", "Video commitments and ARCHITECTURE GATE text", "Drive: 'Aurit Mediation | 2026 Day Strategic Plan - Internal', read 2026-09-18"],
          ["YouTube RSS / oEmbed; Vimeo oEmbed; ClickUp task 86b6nuw3n", "Existing video asset inventory (two YouTube channels, 16 Vimeo assets, 3 site embeds)", "Read 2026-09-18"],
          ["Workflow research outputs (7 lenses), synthesis, adversarial verdicts, legal verification, completeness critic", "Multi-agent live research run", "2026-09-18; JSON in deliverable folder"]]
    sheet(wb, "Evidence Sources", "Evidence behind this calendar", ["Files live in the deliverable's evidence folder in the SEO_Dept repository."],
          ["Source", "What it contains", "Date / provenance"], ev, [48, 80, 50])
    wb.save(out); print("saved", out)

if __name__ == "__main__":
    main(*sys.argv[1:4])
