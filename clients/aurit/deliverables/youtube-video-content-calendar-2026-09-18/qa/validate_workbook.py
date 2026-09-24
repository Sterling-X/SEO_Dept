#!/usr/bin/env python3
"""Deterministic checks on the built workbook. Usage: validate_workbook.py out.xlsx evidence/aurit_live_pages.csv
Exit 1 on any FAIL."""
import sys, csv, re
from openpyxl import load_workbook
from urllib.parse import urlparse
wb = load_workbook(sys.argv[1]); live = {r["path"].rstrip("/") + "/" for r in csv.DictReader(open(sys.argv[2]))}
fails, warns = [], []
def rows(ws):
    hdr_i = None
    for i, r in enumerate(ws.iter_rows(values_only=True), 1):
        if r and r[0] in ("Order", "Recording Title", "Clip Title / Opening Question", "Candidate Topic", "Area", "Source"): hdr_i = i; hdr = list(r); break
    if hdr_i is None: fails.append(f"{ws.title}: no header row"); return [], []
    data = [dict(zip(hdr, r)) for r in ws.iter_rows(min_row=hdr_i + 1, values_only=True) if any(c not in (None, "") for c in r)]
    return hdr, data
need = ["Legend", "Calendar", "Recording Briefs", "Shorts", "Topic Research", "Publishing Plan", "AI Search Briefs", "Legal Verification", "Evidence Sources"]
for n in need:
    if n not in wb.sheetnames: fails.append(f"missing tab {n}")
_, cal = rows(wb["Calendar"]); _, briefs = rows(wb["Recording Briefs"]); _, shorts = rows(wb["Shorts"]); _, topics = rows(wb["Topic Research"]); _, plan = rows(wb["Publishing Plan"]); _, ai = rows(wb["AI Search Briefs"]); _, legal = rows(wb["Legal Verification"])
orders = [r["Order"] for r in cal]
if orders != sorted(orders) or len(set(orders)) != len(orders): fails.append("Calendar: Order not unique ascending")
titles = [r["Recording Title"] for r in cal]; brief_titles = {b["Recording Title"] for b in briefs}; ai_titles = {a["Recording Title"] for a in ai}
for r in cal:
    t = r["Recording Title"]; new = not str(r["Track / Batch"]).lower().startswith("track 0") and "activate" not in str(r["Track / Batch"]).lower()
    if t not in brief_titles: fails.append(f"Calendar '{t}': no Recording Brief")
    n_shorts = sum(1 for s in shorts if s["Source Video"] == t)
    if new and n_shorts < 2: fails.append(f"Calendar '{t}': {n_shorts} shorts (<2)")
    if new and t not in ai_titles: fails.append(f"Calendar '{t}': no AI Search Brief")
    link = str(r["Internal Link Target"] or "")
    for u in re.split(r"[;\n ]+", link):
        if not u or u.lower().startswith("architecture"): continue
        p = urlparse(u).path if u.startswith("http") else u
        p = (p.rstrip("/") + "/") if p else p
        if p and p not in live: fails.append(f"Calendar '{t}': Internal Link Target not a live page: {u}")
    for k in ("Topic Signal", "Why It Matters Commercially", "Featured Speaker", "CTA / Next Step", "Hub Section"):
        if not r.get(k): fails.append(f"Calendar '{t}': empty {k}")
    if str(r["Notes"]).find("ARCHITECTURE GATE") < 0 and not link: fails.append(f"Calendar '{t}': no link target and no gate flag")
for b in briefs:
    if b["Recording Title"] not in titles: warns.append(f"Brief '{b['Recording Title']}' has no Calendar row")
    p = str(b["Existing Page to Support"] or "")
    pp = urlparse(p).path if p.startswith("http") else p
    if pp and (pp.rstrip("/") + "/") not in live and "gate" not in p.lower(): fails.append(f"Brief '{b['Recording Title']}': Existing Page not live: {p}")
for s in shorts:
    if s["Source Video"] not in titles and not any(k in str(s["Source Video"]).lower() for k in ("faq", "existing", "vimeo", "asset")): warns.append(f"Short '{s['Clip Title / Opening Question']}' source not in Calendar: {s['Source Video']}")
bad = re.compile(r"—|\bTBD\b|lorem|\[placeholder\]|XXX|\bTODO\b|INSERT ", re.I)
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and bad.search(c.value): fails.append(f"{ws.title}!{c.coordinate}: placeholder/em dash: {c.value[:60]!r}")
    if ws.title not in ("Legend",) and not ws.freeze_panes: fails.append(f"{ws.title}: no frozen header")
hdr_font_ok = all(wb[n].cell(row=[i for i,r in enumerate(wb[n].iter_rows(values_only=True),1) if r and r[0] in ("Order","Recording Title","Clip Title / Opening Question","Candidate Topic","Area","Source")][0], column=1).fill.fgColor.rgb.endswith("101F3C") for n in need if n != "Legend")
if not hdr_font_ok: fails.append("header fill not navy 101F3C on every tab")
if len(legal) == 0: warns.append("Legal Verification tab has no rows")
print(f"calendar={len(cal)} briefs={len(briefs)} shorts={len(shorts)} topics={len(topics)} plan={len(plan)} ai={len(ai)} legal={len(legal)}")
for w in warns: print("WARN", w)
for f in fails: print("FAIL", f)
print("RESULT:", "PASS" if not fails else f"FAIL ({len(fails)})"); sys.exit(1 if fails else 0)
